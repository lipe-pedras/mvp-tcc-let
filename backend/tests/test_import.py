import io

import pytest
from docx import Document as Docx

from app.main import app
from app.models.user import Role
from app.services.ingestion.parsers import DoclingParser, get_parser
from tests.conftest import auth


def make_docx() -> bytes:
    d = Docx()
    d.add_heading("Reembolso de viagens", level=1)
    d.add_paragraph("O prazo para solicitar é de 15 dias.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Item", "Teto"
    t.cell(1, 0).text, t.cell(1, 1).text = "Refeição", "R$ 78,50"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


@pytest.fixture
def gestor(groups, make_user):
    return make_user("rh@empresa.com", Role.gestor, groups["todos"])


@pytest.fixture(autouse=True)
def real_parser():
    app.dependency_overrides[get_parser] = DoclingParser
    yield


def test_docx_import_keeps_table_as_markdown_and_starts_as_draft(client, gestor, groups):
    r = client.post(
        "/api/documents/import",
        headers=auth(gestor),
        files={"file": ("reembolso.docx", make_docx())},
        data={"title": "Reembolso", "group_ids": [groups["todos"].id]},
    )
    assert r.status_code == 201, r.text
    doc = r.json()
    assert doc["status"] == "draft"
    assert "15 dias" in doc["content_md"]
    assert "| Item" in doc["content_md"] and "R$ 78,50" in doc["content_md"]


def test_unsupported_extension_rejected(client, gestor, groups):
    r = client.post(
        "/api/documents/import",
        headers=auth(gestor),
        files={"file": ("x.exe", b"MZ")},
        data={"title": "X", "group_ids": [groups["todos"].id]},
    )
    assert r.status_code == 422


def test_collaborator_cannot_import(client, groups, make_user):
    colab = make_user("c@empresa.com", Role.colaborador, groups["todos"])
    r = client.post(
        "/api/documents/import",
        headers=auth(colab),
        files={"file": ("a.docx", make_docx())},
        data={"title": "X", "group_ids": [groups["todos"].id]},
    )
    assert r.status_code == 403
